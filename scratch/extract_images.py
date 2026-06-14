import os
from pptx import Presentation

def main():
    prs = Presentation("PRESENTAZIONE/presentation_v3.pptx")
    out_dir = "scratch/extracted_images"
    os.makedirs(out_dir, exist_ok=True)
    
    for i, slide in enumerate(prs.slides):
        for j, shape in enumerate(slide.shapes):
            if shape.shape_type == 13: # Picture
                image = shape.image
                ext = image.ext
                filename = f"slide_{i+1}_shape_{j}_{shape.name}.{ext}"
                path = os.path.join(out_dir, filename)
                with open(path, "wb") as f:
                    f.write(image.blob)
                print(f"Extracted image to: {path}")

if __name__ == "__main__":
    main()
