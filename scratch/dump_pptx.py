from pptx import Presentation
import sys

prs = Presentation("PRESENTAZIONE/presentation_v5.pptx")
with open("scratch/v5_dump.txt", "w", encoding="utf-8") as f:
    for i, slide in enumerate(prs.slides):
        f.write(f"\n{'='*40}\nSLIDE {i+1}\n{'='*40}\n")
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                f.write(f"[{shape.shape_type}] {shape.text}\n")
            elif shape.has_table:
                f.write("[TABLE]\n")
                for row in shape.table.rows:
                    f.write(" | ".join([cell.text_frame.text.replace('\n', ' ') for cell in row.cells]) + "\n")
