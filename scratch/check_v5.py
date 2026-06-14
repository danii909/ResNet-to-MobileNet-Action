"""
Check if any slides have issues (shapes outside bounds, etc.)
"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pptx import Presentation
from pptx.util import Inches

prs = Presentation(r"PRESENTAZIONE/presentation_v5.pptx")
W = prs.slide_width
H = prs.slide_height
issues = []
for i, slide in enumerate(prs.slides):
    for shape in slide.shapes:
        l, t, w, h = shape.left, shape.top, shape.width, shape.height
        if l + w > W + Inches(0.1):
            issues.append(f"Slide {i+1} shape '{shape.name}': right edge {(l+w)/914400:.2f}in > {W/914400:.2f}in")
        if t + h > H + Inches(0.1):
            issues.append(f"Slide {i+1} shape '{shape.name}': bottom edge {(t+h)/914400:.2f}in > {H/914400:.2f}in")

if issues:
    print("Issues found:")
    for iss in issues:
        print(" ", iss)
else:
    print("No layout issues found! All shapes within bounds.")

# Check images are embedded correctly
print("\nImages in presentation:")
for i, slide in enumerate(prs.slides):
    for shape in slide.shapes:
        if shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
            print(f"  Slide {i+1}: Picture '{shape.name}' {shape.width/914400:.1f}×{shape.height/914400:.1f}in")
