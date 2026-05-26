from pptx import Presentation
from pptx.util import Inches
import os

def main():
    prs = Presentation("PRESENTAZIONE/presentation_v4.pptx")
    out_lines = []
    for i, slide in enumerate(prs.slides):
        out_lines.append(f"\n--- Slide {i+1} ---")
        for j, shape in enumerate(slide.shapes):
            shape_type = shape.shape_type
            name = shape.name
            left = shape.left / Inches(1)
            top = shape.top / Inches(1)
            width = shape.width / Inches(1)
            height = shape.height / Inches(1)
            out_lines.append(f"Shape {j}: Type={shape_type}, Name='{name}', Left={left:.2f}, Top={top:.2f}, Width={width:.2f}, Height={height:.2f}")
            
            if shape.has_text_frame:
                text = " | ".join([p.text for p in shape.text_frame.paragraphs if p.text.strip()])
                if text:
                    out_lines.append(f"  Text: {text}")
            
            if shape_type == 13: # Picture
                out_lines.append(f"  Picture: width={shape.width/Inches(1):.2f}, height={shape.height/Inches(1):.2f}")
                
    with open("scratch/inspect_output_v4.txt", "w", encoding="utf-8") as f:
        f.write("\n".join(out_lines))
    print("Done inspection.")

if __name__ == "__main__":
    main()
