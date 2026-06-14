import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from pptx import Presentation

prs = Presentation(r'PRESENTAZIONE/presentation_v4.pptx')
for i, slide in enumerate(prs.slides):
    print(f"=== Slide {i+1} [{slide.slide_layout.name}] ===")
    for shape in slide.shapes:
        if shape.has_text_frame:
            text = ' | '.join([p.text for p in shape.text_frame.paragraphs if p.text.strip()])
            if text:
                print(f"  [{shape.name}]: {text[:400]}")
    print()

print(f"\nTotal slides: {len(prs.slides)}")
print(f"Slide width: {prs.slide_width.inches:.1f} in")
print(f"Slide height: {prs.slide_height.inches:.1f} in")
