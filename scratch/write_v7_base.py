import os

content = '''"""
presentation_v7.pptx
Generated automatically based on user's v5 structure with schematic updates.
"""
import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# Colors
C_BG        = RGBColor(0xFF, 0xFF, 0xFF)
C_PRIMARY   = RGBColor(0x1A, 0x1A, 0x2E)
C_HIGHLIGHT = RGBColor(0x0F, 0x3D, 0x91)
C_HIGHLIGHT2= RGBColor(0xE5, 0x48, 0x00)
C_GRAY      = RGBColor(0x55, 0x55, 0x55)
C_LGRAY     = RGBColor(0xDD, 0xDD, 0xDD)
C_BOX       = RGBColor(0xF0, 0xF4, 0xFF)
C_BOX2      = RGBColor(0xFF, 0xF3, 0xE0)
C_WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
C_SUCCESS   = RGBColor(0x28, 0xA7, 0x45)
C_ERROR     = RGBColor(0xDC, 0x35, 0x45)

# Paths
BASE   = r"results/Train-eval-test-split (group-aware)"
BENCH  = r"results/Test_Set_Benchmark"
GA     = f"{BASE}/Training"

T20_FIGS = f"{GA}/slurm-train-eval-4565/kd_t20_a07_24f_lightaug/figures"
T20_CM   = f"{BASE}/Confusion-Matrix/kd_t20_a07_24f_lightaug"
T20_TSNE = f"{GA}/slurm-train-eval-4565/tsne_plots"
BL_FIGS  = f"{GA}/slurm-train-eval-4495/baseline_ls005_24f_lightaug/figures"
BL_CM    = f"{BASE}/Confusion-Matrix/baseline_ls005_24f_lightaug"
BL_TSNE  = f"{GA}/slurm-train-eval-4495/tsne_plots"

HIST_TOP1   = f"{BASE}/distillation_histogram_test_top1.png"
HIST_TOP5   = f"{BASE}/distillation_histogram_test_top5.png"
COMP_CHART  = f"{BASE}/distillation_comparison.png"
DELTA_CHART = f"{BASE}/figures/delta_accuracy_kd_vs_baseline.png"

GPU_PLT  = f"{BASE}/figures/gpu_benchmark_plots.png"
CPU_PLT  = f"{BASE}/figures/cpu_benchmark_plots.png"
COMB_PLT = f"{BASE}/figures/combined_benchmark_plots.png"

AT_TSNE_SYM = f"{GA}/slurm-train-eval-4601/tsne_plots/tsne_teacher_baseline_distill_at_symmetric.png"
AT_TSNE_TMP = f"{GA}/slurm-train-eval-4601/tsne_plots/tsne_teacher_baseline_distill_at_temporal.png"

W = Inches(13.33)
H = Inches(7.5)

prs = Presentation()
prs.slide_width  = W
prs.slide_height = H
TOTAL_SLIDES = 22

def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])

def set_bg(slide, color=C_BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_rect(slide, left, top, width, height, fill_color=None, line_color=None):
    shape = slide.shapes.add_shape(1, left, top, width, height)
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    return shape

def add_textbox(slide, text, left, top, width, height,
                font_size=12, bold=False, color=C_PRIMARY,
                align=PP_ALIGN.LEFT, italic=False, font_name="Calibri"):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font_name
    return tb

def add_multiline(slide, lines, left, top, width, height,
                  font_size=11, bold=False, color=C_PRIMARY,
                  align=PP_ALIGN.LEFT, font_name="Calibri", line_spacing=None):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for item in lines:
        if isinstance(item, str):
            txt, b, c, fs = item, bold, color, font_size
        else:
            txt, b, c, fs = item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        if line_spacing:
            p.line_spacing = Pt(line_spacing)
        run = p.add_run()
        run.text = txt
        run.font.size = Pt(fs)
        run.font.bold = b
        run.font.color.rgb = c
        run.font.name = font_name
    return tb

def footer(slide, n, total=TOTAL_SLIDES):
    add_textbox(slide, f"{n}/{total}",
                Inches(11.8), Inches(7.05), Inches(1.3), Inches(0.35),
                font_size=10, color=C_GRAY, align=PP_ALIGN.RIGHT)

def section_header(slide, text, color=C_HIGHLIGHT):
    add_rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.08), fill_color=color)
    add_textbox(slide, text, Inches(0.45), Inches(0.08), Inches(10), Inches(0.42),
                font_size=9, bold=True, color=color)

def slide_title(slide, title, subtitle=""):
    add_textbox(slide, title, Inches(0.45), Inches(0.55), Inches(12.4), Inches(0.7),
                font_size=26, bold=True, color=C_PRIMARY)
    if subtitle:
        add_textbox(slide, subtitle, Inches(0.45), Inches(1.15), Inches(12.4), Inches(0.4),
                    font_size=12, color=C_GRAY, italic=True)
        add_rect(slide, Inches(0.45), Inches(1.5), Inches(12.43), Inches(0.02), fill_color=C_LGRAY)
    else:
        add_rect(slide, Inches(0.45), Inches(1.2), Inches(12.43), Inches(0.02), fill_color=C_LGRAY)

def pic(slide, path, left, top, width, height):
    try:
        return slide.shapes.add_picture(path, left, top, width, height)
    except:
        return add_rect(slide, left, top, width, height, fill_color=C_LGRAY)
'''

with open("scratch/make_presentation_v7.py", "w", encoding="utf-8") as f:
    f.write(content)
