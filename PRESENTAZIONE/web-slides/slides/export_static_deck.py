from __future__ import annotations

import argparse
import re
import subprocess
import tempfile
from pathlib import Path

from PIL import Image, JpegImagePlugin
from pptx import Presentation
from pptx.util import Inches


SLIDE_W = 1600
SLIDE_H = 900


def chrome_path() -> Path:
    candidates = [
        Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe"),
        Path(r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe"),
        Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
        Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    raise FileNotFoundError("Chrome/Edge executable not found in standard locations.")


def count_slides(html: str) -> int:
    # Remove HTML comments to avoid counting commented-out/inactive slides
    clean_html = re.sub(r"<!--.*?-->", "", html, flags=re.DOTALL)
    return len(re.findall(r'<div\s+class="[^"]*\bslide\b', clean_html))


def html_for_slide(html: str, slide_number: int, base_href: str) -> str:
    base = f'  <base href="{base_href}">\n'
    override = f"""
  <style id="static-export-override">
    html, body {{
      width: {SLIDE_W}px !important;
      height: {SLIDE_H}px !important;
      margin: 0 !important;
      overflow: hidden !important;
      background: var(--white) !important;
    }}
    #deck {{
      width: {SLIDE_W}px !important;
      height: {SLIDE_H}px !important;
      overflow: hidden !important;
    }}
    #nav, #image-modal {{
      display: none !important;
    }}
    #deck > .slide {{
      opacity: 0 !important;
      pointer-events: none !important;
      transition: none !important;
      transform: translate(-50%, -50%) scale(1) !important;
      z-index: 0 !important;
    }}
    #deck > .slide:nth-child({slide_number}) {{
      opacity: 1 !important;
      pointer-events: auto !important;
      z-index: 10 !important;
    }}
  </style>
"""
    return html.replace("</head>", base + override + "\n</head>")


def render_pngs(source: Path, out_dir: Path) -> list[Path]:
    html = source.read_text(encoding="utf-8")
    base_href = source.parent.as_uri() + "/"
    total = count_slides(html)
    if total == 0:
        raise RuntimeError("No slides found.")

    out_dir.mkdir(parents=True, exist_ok=True)
    browser = chrome_path()
    pngs: list[Path] = []

    with tempfile.TemporaryDirectory(prefix="static-deck-") as tmp:
        tmp_dir = Path(tmp)
        for idx in range(1, total + 1):
            tmp_html = tmp_dir / f"slide-{idx:02d}.html"
            png = out_dir / f"slide-{idx:02d}.png"
            tmp_html.write_text(html_for_slide(html, idx, base_href), encoding="utf-8")

            subprocess.run(
                [
                    str(browser),
                    "--headless=new",
                    "--disable-gpu",
                    "--hide-scrollbars",
                    "--allow-file-access-from-files",
                    f"--window-size={SLIDE_W},{SLIDE_H}",
                    "--force-device-scale-factor=1",
                    f"--screenshot={png}",
                    tmp_html.as_uri(),
                ],
                check=True,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
            )
            pngs.append(png)
            print(f"Rendered {png.name}")

    return pngs


def make_pdf(pngs: list[Path], output: Path) -> None:
    images = [Image.open(p).convert("RGB") for p in pngs]
    first, rest = images[0], images[1:]
    first.save(output, save_all=True, append_images=rest, resolution=144)
    for image in images:
        image.close()


def make_pptx(pngs: list[Path], output: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)
    blank = prs.slide_layouts[6]

    for png in pngs:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(png), 0, 0, width=prs.slide_width, height=prs.slide_height)

    prs.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("--out", type=Path, default=Path("exports"))
    args = parser.parse_args()

    source = args.source.resolve()
    out = args.out.resolve()
    stem = source.stem
    image_dir = out / f"{stem}_png"
    pdf = out / f"{stem}_static.pdf"
    pptx = out / f"{stem}_static.pptx"

    pngs = render_pngs(source, image_dir)
    make_pdf(pngs, pdf)
    make_pptx(pngs, pptx)

    print(f"PDF: {pdf}")
    print(f"PPTX: {pptx}")
    print(f"PNGs: {image_dir}")


if __name__ == "__main__":
    main()
